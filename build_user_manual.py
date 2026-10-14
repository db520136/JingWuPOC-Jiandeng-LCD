from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = "设备用户操作说明手册.docx"
FONT = "Microsoft YaHei"
ACCENT = "164E63"
ACCENT_LIGHT = "E8F3F5"
TABLE_HEADER = "1F4E5F"
TABLE_ALT = "F4F8F9"
BORDER = "D9E1E5"
GRAY = "5B6770"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=130, bottom=100, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_font(paragraph, size=10.5, color="000000"):
    for run in paragraph.runs:
        set_run_font(run, size=size, color=color)


def style_document(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18

    for name, size, space_before, space_after in (
        ("Title", 26, 0, 16),
        ("Heading 1", 18, 14, 8),
        ("Heading 2", 13, 10, 5),
        ("Heading 3", 11, 7, 3),
    ):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(space_after)
        style.paragraph_format.keep_with_next = True

    if "Small Note" not in styles:
        note = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        note = styles["Small Note"]
    note.font.name = FONT
    note._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    note.font.size = Pt(9)
    note.font.color.rgb = RGBColor.from_string(GRAY)
    note.paragraph_format.space_after = Pt(4)
    note.paragraph_format.line_spacing = 1.1


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = "设备用户操作说明手册  |  以设备实际显示为准"
    for run in p.runs:
        set_run_font(run, size=8.5, color=GRAY)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style="Heading " + str(level))
    p.add_run(text)
    return p


def add_para(doc, text="", style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r)
        p.paragraph_format.space_after = Pt(3)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(item)
        set_run_font(r)
        p.paragraph_format.space_after = Pt(3)


def add_table(doc, headers, rows, widths=None, compact=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, TABLE_HEADER)
        set_cell_border(cell)
        set_cell_margins(cell, 110, 130, 110, 130)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(header))
        set_run_font(r, size=9.5 if compact else 10, bold=True, color="FFFFFF")
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for cidx, value in enumerate(row):
            cell = cells[cidx]
            if ridx % 2 == 1:
                set_cell_shading(cell, TABLE_ALT)
            set_cell_border(cell)
            set_cell_margins(cell, 100, 130, 100, 130)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if cidx else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_run_font(r, size=9 if compact else 9.5)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_note(doc, title, text):
    p = doc.add_paragraph(style="Small Note")
    r = p.add_run(title + "：")
    set_run_font(r, size=9, bold=True, color=ACCENT)
    r = p.add_run(text)
    set_run_font(r, size=9, color=GRAY)
    return p


def page_break(doc):
    doc.add_page_break()


def build():
    doc = Document()
    style_document(doc)
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)
    add_footer(sec)

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("设备用户操作说明手册")
    set_run_font(r, size=27, bold=True, color="000000")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("警务肩灯一体设备")
    set_run_font(r, size=15, color=ACCENT)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("面向设备使用人员的操作、状态查看与故障处理说明")
    set_run_font(r, size=11, color=GRAY)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(115)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("适用于当前项目版本")
    set_run_font(r, size=10.5, color=GRAY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("请以设备实际显示文字和图标为准")
    set_run_font(r, size=9.5, color=GRAY)

    page_break(doc)

    add_heading(doc, "使用前须知", 1)
    add_para(doc, "本手册说明设备从开机、登录、定位、对讲到设置和升级的日常使用方法。设备会在后台处理网络重连、定位采集、电量刷新和升级状态，使用人员只需根据屏幕提示操作。")
    add_para(doc, "涉及服务器账号、组呼权限、单呼成员和升级任务的功能，需要网络连接正常并由平台下发相应权限。游客模式用于查看定位和基础状态，不具备对讲权限。")
    add_note(doc, "重要", "首次使用前请确认设备电量充足、SIM 卡已插入，并在信号良好的位置完成首次登录。升级过程中不要关机或拔出电源。")

    add_heading(doc, "目录", 1)
    toc = [
        "1 认识设备界面和状态图标",
        "2 开机与登录",
        "3 游客模式",
        "4 定位功能",
        "5 组呼和单呼",
        "6 PTT 对讲",
        "7 肩灯与双击模式",
        "8 音量操作",
        "9 屏幕亮度和自动息屏",
        "10 电量和充电状态",
        "11 网络状态和 SIM 卡信息",
        "12 软件升级",
        "13 用户切换和强制下线",
        "14 常见提示和处理方法",
        "15 用户操作速查表",
    ]
    add_numbered(doc, toc)

    page_break(doc)

    add_heading(doc, "1 认识设备界面和状态图标", 1)
    add_para(doc, "设备登录或进入游客模式后，屏幕底部有四个主功能入口：定位、肩灯、对讲机和设置。点击图标或对应区域进入功能页；子页面通常通过返回手势或系统页面导航返回上一级。")
    add_table(doc, ["界面", "主要内容", "可执行操作"], [
        ("定位", "时间、星期、姓名或游客、经纬度、海拔、卫星数量", "进入定位详情，查看当前账号并切换用户"),
        ("肩灯", "肩灯模式、双击模式", "选择关闭、红蓝交替闪或红蓝对角闪；设置双击功能"),
        ("对讲机", "组呼模式、单呼模式", "登录且有权限时选择组呼小组或在线单呼成员"),
        ("设置", "设备信息、亮度、音量、息屏、网络、跌倒检测、升级", "查看状态或修改设备参数"),
    ], widths=[2.2, 7.0, 7.5])
    add_heading(doc, "状态栏", 2)
    add_table(doc, ["显示内容", "含义"], [
        ("警号或游客", "显示当前登录身份；游客模式显示“游客”"),
        ("信号图标", "网络和 TCP 链路均正常时按信号等级显示；显示最低状态时表示链路尚未完整建立"),
        ("电量", "显示当前电量百分比；充电时电量文字显示为绿色"),
    ], widths=[4.0, 12.7])
    add_note(doc, "提示", "网络断开或后台重连通常不会阻止定位页面、游客模式或自动息屏。需要对讲时，设备会根据当前连接状态给出提示。")

    add_heading(doc, "2 开机与登录", 1)
    add_heading(doc, "2.1 开机", 2)
    add_numbered(doc, [
        "按设备电源键开机，等待启动页面完成。",
        "启动完成后进入登录页面。登录页面包含 M、F 前缀选择、六位数字输入框、删除和确认按键。",
        "开机后设备会在后台尝试建立网络和 TCP 连接；连接提示只在登录页生命周期内按需显示，后续重连主要在后台执行。",
    ])
    add_heading(doc, "2.2 普通警号登录", 2)
    add_numbered(doc, [
        "根据警号首字母选择 M 或 F。",
        "依次输入六位数字警号。输入错误时点击“删除”修改。",
        "点击“确认”。屏幕显示登录进度，设备等待服务器应答。",
        "登录成功后进入定位一级页面，并刷新当前警员姓名和可用组呼、单呼信息。",
    ])
    add_table(doc, ["提示", "说明"], [
        ("请输入6位编号", "输入的数字不足六位，请补齐后再确认"),
        ("等待网络", "设备正在等待蜂窝网络准备完成"),
        ("登录中", "已提交登录请求，正在等待服务器应答"),
        ("警号未注册或登录失败", "服务器已明确拒绝本次登录，请核对警号或联系管理员"),
        ("网络连接中或 TCP 连接中", "当前连接尚未完成，可等待恢复后再次确认登录"),
    ], widths=[5.0, 11.7])
    add_heading(doc, "2.3 无网络时的处理", 2)
    add_para(doc, "如果确认登录时网络或 TCP 尚未准备好，设备会弹出网络状态确认框。选择继续后可进入游客模式；取消则留在登录页，后台继续尝试连接。")
    add_note(doc, "登录结果", "只有本次登录事务收到成功应答后，设备才会进入普通登录业务页。被拒绝的登录不会被当作成功。")

    add_heading(doc, "3 游客模式", 1)
    add_para(doc, "在登录页面输入 000000 并点击“确认”，设备直接进入本地游客模式，不发送 TCP 登录包。游客身份在状态栏和定位详情中显示为“游客”。")
    add_table(doc, ["游客模式支持", "游客模式不支持"], [
        ("查看定位一级页面和定位详情", "组呼、单呼和 PTT 对讲"),
        ("查看电量、信号和设备状态", "加入组呼小组或选择单呼成员"),
        ("使用肩灯、音量、亮度、息屏和部分设置", "依赖 POC 权限的通信业务"),
    ], widths=[8.2, 8.5])
    add_para(doc, "游客模式仍可后台获取定位并在满足条件时上传定位数据。网络恢复后，设备可能提示“网络正常，是否重新登录？”，确认后返回登录页重新输入警号；取消则继续保持游客身份。")

    add_heading(doc, "4 定位功能", 1)
    add_heading(doc, "4.1 定位一级页面", 2)
    add_para(doc, "定位一级页面显示当前定位摘要，包括姓名或游客、经度、纬度、海拔、卫星数量、时间和星期。点击定位图标进入定位详情。")
    add_heading(doc, "4.2 定位详情", 2)
    add_para(doc, "定位详情显示当前用户姓名和警号。普通登录显示实际警号，游客模式显示“警号：游客”。点击“切换用户”可以结束当前身份并返回登录页面。")
    add_heading(doc, "4.3 定位采集和上报", 2)
    add_bullets(doc, [
        "GNSS 初始化失败时，设备会自动重试，不需要用户反复操作。",
        "初始化成功后先进行短周期采集，最多尝试十次；获得有效定位后可提前进入正常采集周期。",
        "正常采集周期为约 30 秒一次。定位结果满足上报条件时通过 TCP 上传。",
        "对讲业务真正占用通信流程时，定位上报可能暂存，业务结束后再发送；空闲状态和没有组呼时不会被误判为对讲忙碌。",
    ])

    add_heading(doc, "5 组呼和单呼", 1)
    add_note(doc, "使用条件", "组呼和单呼仅对正常登录且具备 POC 权限的账号开放。游客模式或 POC 权限关闭时，相关页面不可用。")
    add_heading(doc, "5.1 选择组呼小组", 2)
    add_numbered(doc, [
        "进入“对讲机”，选择“组呼模式”。",
        "在列表中选择目标小组。列表包含“无/退出”选项；选择其他小组时按屏幕提示确认。",
        "确认后等待入组完成。入组成功后可使用 PTT 对讲。",
    ])
    add_heading(doc, "5.2 选择单呼成员", 2)
    add_numbered(doc, [
        "进入“对讲机”，选择“单呼模式”。",
        "列表只显示当前在线且可呼叫的成员，不显示离线成员和本机。",
        "选择成员并确认。结束单呼时选择“无/解散”。",
    ])
    add_note(doc, "刷新名单", "成员和小组信息由平台名单提供。进入登录后设备会刷新一次，页面中的刷新操作可能需要等待网络请求完成。")

    add_heading(doc, "6 PTT 对讲", 1)
    add_para(doc, "PTT 是按住讲话、松开结束的对讲方式。对讲前请确认已经正常登录并加入可用组呼小组，或已选择单呼成员。")
    add_numbered(doc, [
        "按住对讲键，设备申请麦权。屏幕提示“正在申请麦权”时请保持按住。",
        "获得麦权后，屏幕提示“正在讲话中”，此时对着麦克风讲话。",
        "松开对讲键，设备停止本地采音并释放麦权。",
        "收到远端讲话时，设备自动播放语音，并显示当前讲话成员和通话类型。",
    ])
    add_table(doc, ["现象", "处理"], [
        ("游客身份，无法对讲", "当前是游客模式，需要返回登录页使用已注册警号登录"),
        ("正在申请麦权时间较长", "保持按键并确认信号图标正常；网络未恢复时先松开，待连接正常后重试"),
        ("抢麦失败或释放麦权失败", "松开按键，等待提示消失后重新尝试；持续失败时检查网络状态"),
        ("对讲中定位暂未上报", "属于业务优先级处理，通话结束后会恢复定位上报"),
    ], widths=[6.2, 10.5])

    add_heading(doc, "7 肩灯与双击模式", 1)
    add_heading(doc, "7.1 屏幕设置肩灯模式", 2)
    add_numbered(doc, [
        "进入“肩灯”，选择“肩灯模式”。",
        "选择“关闭”“红蓝交替闪”或“红蓝对角闪”。",
        "出现“确认修改？”后点击“确认”；点击“取消”保持原设置。",
    ])
    add_heading(doc, "7.2 肩灯按键操作", 2)
    add_table(doc, ["操作", "结果"], [
        ("短按，2 秒内松开", "肩灯开启时关闭；肩灯关闭时打开红蓝交替闪；红蓝对角闪时关闭"),
        ("长按至少 2 秒后松开", "切换双击模式开启或关闭，并提示切换后的状态约 2 秒"),
        ("长按达到 2 秒仍不松开", "只触发本次切换，不会每隔 2 秒重复切换；松开并完成防抖后才能再次操作"),
    ], widths=[6.0, 10.7])
    add_para(doc, "双击模式开启时，设备可根据加速度传感器的双击事件切换肩灯；关闭时不响应双击切换，但不影响肩灯按键短按。")

    add_heading(doc, "8 音量操作", 1)
    add_para(doc, "音量可以在“设置 > 音量”页面通过滑块调节，也可以使用设备侧面的音量增加键和音量减少键调节。音量范围为 0 到 11 级。")
    add_bullets(doc, [
        "按下音量键后会尽快播放按键提示音。必须松开后，下一次按下才算新的调节操作。",
        "屏幕会短暂显示音量滑条和当前级数。",
        "音量连续 1 分钟没有变化后自动保存，设备重启后沿用已保存值。",
    ])

    add_heading(doc, "9 屏幕亮度和自动息屏", 1)
    add_heading(doc, "9.1 调节屏幕亮度", 2)
    add_para(doc, "进入“设置 > 屏幕亮度”，拖动竖向滑块调节亮度。界面显示范围为 5% 到 100%。设备对硬件最大亮度做了限制，显示 100% 时已经是当前设备可用的最高体验亮度。")
    add_heading(doc, "9.2 设置自动息屏", 2)
    add_para(doc, "进入“设置 > 息屏时间”，可选择 10 秒、15 秒、20 秒、30 秒、1 分钟或“永不息屏”。选择后按提示确认。")
    add_bullets(doc, [
        "登录成功进入业务页或进入游客模式后，才开始计算息屏时间。",
        "触摸、按键、对讲等活动会重新计时；真正处于对讲忙碌状态时不会息屏。",
        "网络断开、网络重连或后台定位处理不会单独阻止息屏。",
        "屏幕息屏后，按屏幕控制键或触发有效操作可唤醒屏幕。",
        "屏幕亮度连续 1 分钟没有变化后自动保存。",
    ])

    add_heading(doc, "10 电量和充电状态", 1)
    add_para(doc, "状态栏显示电量百分比。电池采样由后台低优先级任务完成，不会阻塞触摸、通信和对讲。")
    add_table(doc, ["显示状态", "含义"], [
        ("绿色电量文字", "设备正在充电"),
        ("绿色 100%", "充电计时达到显示上限，但充电芯片尚未确认充满"),
        ("白色 100%", "充电芯片确认已充满"),
        ("电量最低显示 1%", "电压很低时仍显示 1%，不显示 0%"),
    ], widths=[5.0, 11.7])
    add_note(doc, "使用建议", "长时间存放或低电量使用前请及时充电。电量百分比经过滤波和状态判断，插拔电源、按键或音频负载瞬间变化时，显示不会立即跟随瞬态跳变。")

    add_heading(doc, "11 网络状态和 SIM 卡信息", 1)
    add_para(doc, "进入“设置 > 网络状态”可查看 SIM 卡号、信号值和网络状态，点击“刷新”重新读取。网络状态页面显示的是蜂窝网络状态；状态栏信号图标还会同时考虑 TCP 链路是否建立。")
    add_table(doc, ["项目", "正常显示", "异常显示"], [
        ("SIM 卡号", "显示 ICCID", "显示“--”，表示暂未读取到"),
        ("信号值", "显示信号等级和 dBm", "显示“无”或信号等级较低"),
        ("网络状态", "成功连接", "连接失败"),
    ], widths=[4.0, 6.0, 6.7])
    add_note(doc, "排查顺序", "遇到登录或对讲异常时，先确认 SIM 卡已插入，再查看信号和网络状态，最后重新尝试登录或 PTT。")

    add_heading(doc, "12 软件升级", 1)
    add_numbered(doc, [
        "进入“设置 > 升级”。",
        "查看当前软件版本，点击“查询新版本”。",
        "有可用版本时，按屏幕提示确认下载和升级。设备会显示下载、准备升级和重启等状态。",
        "设备重启完成后再次进入升级页面，核对当前软件版本。",
    ])
    add_table(doc, ["升级阶段", "用户注意事项"], [
        ("查询和下载", "保持网络稳定，不要反复点击查询按钮"),
        ("准备升级", "保持设备供电，等待设备自行处理"),
        ("重启升级", "不要断电、拔电或操作按键"),
        ("失败提示", "记录提示内容，确认网络和电量后再按提示重试；以重启后版本查询结果为最终判断"),
    ], widths=[5.0, 11.7])

    add_heading(doc, "13 用户切换和强制下线", 1)
    add_heading(doc, "13.1 主动切换用户", 2)
    add_numbered(doc, [
        "在定位详情页面点击“切换用户”。",
        "在确认框中点击“确认”，屏幕显示“正在退出... ...”。旧账号的对讲操作会立即屏蔽。",
        "退出流程完成或连接无法恢复后，设备返回登录页面。",
        "重新选择 M 或 F，输入新警号并点击“确认”。游客模式切换用户时直接返回登录页面。",
    ])
    add_heading(doc, "13.2 服务器强制下线", 2)
    add_para(doc, "当同一警号在其他设备上强制登录，当前设备会被服务器退出。设备返回登录页面后播放“账号强制下线”语音提示。听到提示后，请重新输入警号登录。")

    add_heading(doc, "14 常见提示和处理方法", 1)
    add_table(doc, ["现象或提示", "建议处理"], [
        ("登录页面长时间等待网络", "等待后台重连；确认登录时若仍无网络，按提示选择游客模式或稍后重试"),
        ("警号未注册", "核对 M/F 前缀和六位数字，确认账号已在平台注册"),
        ("游客身份，无法对讲", "游客模式不支持对讲，请返回登录页使用正常警号登录"),
        ("组呼或单呼列表为空", "确认已正常登录且有 POC 权限，等待成员名单刷新"),
        ("单呼成员不见了", "三级页面只显示在线成员，成员离线后会从列表隐藏"),
        ("信号图标显示最低状态", "检查网络状态页面，确认 SIM 卡和信号；TCP 未建立时不会显示正常信号等级"),
        ("定位暂时没有有效坐标", "将设备移到室外或开阔位置，等待 GNSS 自动重试"),
        ("升级失败提示", "保持供电和网络，重启后进入升级页面核对版本，再决定是否重试"),
        ("肩灯按键没有立即切换", "确认按键已松开并完成防抖；长按达到 2 秒只切换一次"),
    ], widths=[6.4, 10.3], compact=True)
    add_heading(doc, "需要联系维护人员的情况", 2)
    add_bullets(doc, [
        "开机后持续无法进入登录页面。",
        "多次确认网络、SIM 卡和供电正常后仍无法建立连接。",
        "I2C 肩灯、音频、屏幕或 GNSS 出现持续硬件错误。",
        "升级重启后版本未变化且重复失败。",
    ])

    add_heading(doc, "15 用户操作速查表", 1)
    add_table(doc, ["目标", "操作路径"], [
        ("正常登录", "登录页选择 M/F > 输入六位警号 > 确认"),
        ("游客登录", "登录页输入 000000 > 确认"),
        ("进入定位详情", "定位一级页面 > 点击定位图标"),
        ("选择组呼", "对讲机 > 组呼模式 > 选择小组 > 确认"),
        ("选择单呼", "对讲机 > 单呼模式 > 选择在线成员 > 确认"),
        ("发起讲话", "按住对讲键申请麦权 > 获得麦权后讲话 > 松开结束"),
        ("切换肩灯", "肩灯 > 肩灯模式 > 选择模式 > 确认"),
        ("切换双击模式", "肩灯 > 双击模式 > 开启或关闭 > 确认；或长按肩灯按键至少 2 秒"),
        ("调节亮度", "设置 > 屏幕亮度或肩灯亮度 > 拖动滑块"),
        ("设置息屏", "设置 > 息屏时间 > 选择时间 > 确认"),
        ("查看网络", "设置 > 网络状态 > 查看 SIM 卡、信号和网络状态 > 刷新"),
        ("升级软件", "设置 > 升级 > 查询新版本 > 按提示下载和重启"),
        ("切换用户", "定位详情 > 切换用户 > 确认 > 返回登录页"),
    ], widths=[5.0, 11.7], compact=True)
    add_note(doc, "版本说明", "本手册按当前项目版本整理。界面文字、可用选项和服务器提示可能随平台配置变化；实际操作时请优先遵循设备当前屏幕提示。")

    # Ensure all paragraphs use the selected East Asian font and avoid widows where possible.
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True
        if paragraph.style.name not in ("Title", "Heading 1", "Heading 2", "Heading 3"):
            set_paragraph_font(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.widow_control = True

    doc.core_properties.title = "设备用户操作说明手册"
    doc.core_properties.subject = "警务肩灯一体设备用户操作、状态查看与故障处理"
    doc.core_properties.author = ""
    doc.save(OUT)


if __name__ == "__main__":
    build()
